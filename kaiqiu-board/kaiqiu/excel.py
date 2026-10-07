"""由当前快照生成 Excel 报表。

所有增长率都由程序算好后写成数值（不依赖 Excel 公式重算），
所以在 WPS、Numbers、预览或网页版 Excel 打开都是同样的数字。
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .source import NATIONAL

NAVY, BAND, WHITE = "17365D", "EEF3FA", "FFFFFF"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color=WHITE)
INT, PCT = "#,##0", "+0.0%;-0.0%;0.0%"
METRICS = {"participants": "参赛人次", "events": "比赛场次"}


def _rate(cur, base):
    return None if base in (None, 0) or cur is None else (cur - base) / base


def _streak(values, i):
    n = 0
    while i > 0 and values[i] > values[i - 1]:
        n, i = n + 1, i - 1
    return n


def last_complete_index(snap) -> int:
    idx = [i for i, m in enumerate(snap["months"]) if m < snap["current_month"]]
    return idx[-1] if idx else len(snap["months"]) - 1


def _table(ws, top: int, headers: list[str], rows: list[list], formats: dict[int, str] | None = None,
           widths: dict[int, float] | None = None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(top, c, h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, top + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            if formats and c in formats and isinstance(v, (int, float)):
                cell.number_format = formats[c]
    ws.freeze_panes = ws.cell(top + 1, 2)
    if rows:
        ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{top + len(rows)}"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 12)


def _title(ws, text: str, sub: str):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=15, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(color="666666")


def build(snap: dict) -> bytes:
    months, data, areas = snap["months"], snap["data"], snap["areas"]
    regions = [a for a in areas if a != NATIONAL]
    cur = snap["current_month"]
    li = last_complete_index(snap)
    lm = months[li]
    fetched = snap["fetched_at"][:16].replace("T", " ")
    wb = Workbook()

    # 1 说明
    ws = wb.active
    ws.title = "说明"
    _title(ws, "开球网 全国及35区域 月度统计", f"数据取得于 {fetched}（北京时间）· 由开球网数据看板 {__version__} 生成")
    notes = [
        ("统计月份", f"{months[0]} 至 {months[-1]}，共 {len(months)} 个月；最近完整月 {lm}"),
        ("进行中月份", f"{cur} 尚未结束，数值仍会变化，不参与排名" if months[-1] == cur else "无"),
        ("参赛人次", "按来源累计参与次数，一人参加多场会重复计入，不是去重人数"),
        ("比赛场次", "来源字段“本月比赛场次”，与比赛盘数不同"),
        ("环比 / 同比", "与上月 / 去年同月相比；基数为 0 或缺失时留空，不当作 0"),
        ("连续增长", "从该月向前，每月严格大于前月的连续次数"),
        ("全国 vs 区域合计", "全国来自独立接口，可能与 35 区域之和不完全相等，两者都原样保留"),
        ("数值说明", "本表所有增长率都是程序计算好的数值，不依赖公式重算"),
        ("数据来源", snap.get("sources", {}).get("monthly", "https://kaiqiuwang.cc/home/cityEventsMonthly.php")),
        ("快照版本", snap["revision"]),
    ]
    for r, (k, v) in enumerate(notes, 4):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 90

    # 2 全国月度
    ws = wb.create_sheet("全国月度")
    _title(ws, "全国月度历史", "环比、同比按各自指标独立计算")
    p, e, o = data[NATIONAL]["participants"], data[NATIONAL]["events"], data[NATIONAL]["over64"]
    rows = []
    for i, m in enumerate(months):
        rows.append([m, p[i], p[i] - p[i - 1] if i else None, _rate(p[i], p[i - 1] if i else None),
                     _rate(p[i], p[i - 12] if i >= 12 else None), _streak(p, i),
                     e[i], _rate(e[i], e[i - 1] if i else None), _rate(e[i], e[i - 12] if i >= 12 else None),
                     o[i], "完整" if m < cur else "进行中"])
    _table(ws, 4, ["月份", "参赛人次", "人次环比增量", "人次环比", "人次同比", "人次连续增长(月)",
                   "比赛场次", "场次环比", "场次同比", "超过64人比赛", "状态"], rows,
           {2: INT, 3: INT, 4: PCT, 5: PCT, 7: INT, 8: PCT, 9: PCT, 10: INT}, {1: 10})
    chart = LineChart()
    chart.title, chart.height, chart.width = "全国参赛人次（月）", 9, 26
    chart.y_axis.title, chart.legend = "人次", None
    chart.add_data(Reference(ws, min_col=2, min_row=4, max_row=4 + len(rows)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(rows)))
    chart.x_axis.tickLblSkip = 12
    ws.add_chart(chart, "M4")

    # 3 区域排名（最近完整月）
    ws = wb.create_sheet("区域排名")
    _title(ws, f"{lm} 区域排名", "按参赛人次降序；不含进行中月份")
    nat = p[li]
    rank = []
    for a in regions:
        v = data[a]["participants"]
        ev = data[a]["events"]
        rank.append([a, v[li], v[li] / nat if nat else None, _rate(v[li], v[li - 1] if li else None),
                     _rate(v[li], v[li - 12] if li >= 12 else None), _streak(v, li), ev[li],
                     _rate(ev[li], ev[li - 12] if li >= 12 else None)])
    rank.sort(key=lambda r: -r[1])
    rows = [[i + 1] + r for i, r in enumerate(rank)]
    _table(ws, 4, ["排名", "区域", "参赛人次", "占全国", "环比", "同比", "连续增长(月)", "比赛场次", "场次同比"],
           rows, {3: INT, 4: "0.0%", 5: PCT, 6: PCT, 8: INT, 9: PCT}, {1: 7})
    chart = BarChart()
    chart.type, chart.title, chart.height, chart.width = "bar", f"{lm} 参赛人次前15区域", 10, 18
    chart.legend = None
    chart.add_data(Reference(ws, min_col=3, min_row=4, max_row=19), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=5, max_row=19))
    chart.y_axis.scaling.orientation = "minMax"
    chart.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(chart, "K4")

    # 4/5 宽表
    for key, label in METRICS.items():
        ws = wb.create_sheet(f"{label}宽表")
        _title(ws, f"{label} · 月份 × 区域", "每行一个月，每列一个区域")
        rows = [[m] + [data[a][key][i] for a in areas] for i, m in enumerate(months)]
        _table(ws, 4, ["月份"] + areas, rows, {c: INT for c in range(2, len(areas) + 2)},
               {c: 9 for c in range(2, len(areas) + 2)})

    # 6 年度汇总
    ws = wb.create_sheet("年度汇总")
    years = sorted({m[:4] for m in months})
    _title(ws, "年度参赛人次", "只计完整月份；当年未结束时“计入月数”小于12")
    rows = []
    for y in years:
        idx = [i for i, m in enumerate(months) if m[:4] == y and m < cur]
        if not idx:
            continue
        rows.append([int(y), len(idx)] + [sum(data[a]["participants"][i] for i in idx) for a in areas])
    _table(ws, 4, ["年份", "计入月数"] + areas, rows, {c: INT for c in range(3, len(areas) + 3)},
           {1: 8, 2: 9, **{c: 10 for c in range(3, len(areas) + 3)}})
    chart = BarChart()
    chart.title, chart.height, chart.width, chart.legend = "全国年度参赛人次", 8, 18, None
    chart.add_data(Reference(ws, min_col=3, min_row=4, max_row=4 + len(rows)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(rows)))
    ws.add_chart(chart, f"A{len(rows) + 7}")

    # 7 时段统计
    ws = wb.create_sheet("时段统计")
    _title(ws, "滚动时段区域统计", "近7/30/90/365天，按参赛人次排序；不是自然月")
    rows = []
    for label, items in snap["periods"].items():
        for i, r in enumerate(sorted(items, key=lambda r: -r["participants"]), 1):
            rows.append([label, i, r["region"], r["participants"], r["events"], r["games"]])
    _table(ws, 4, ["时段", "排名", "区域", "参赛人次", "比赛场次", "比赛盘数"], rows, {4: INT, 5: INT, 6: INT})

    # 8 区域月度明细（长表，便于数据透视）
    ws = wb.create_sheet("区域月度明细")
    _title(ws, "全国及各区域月度明细", "长表格式，适合数据透视表")
    rows = []
    for a in areas:
        v, ev, ov = data[a]["participants"], data[a]["events"], data[a]["over64"]
        for i, m in enumerate(months):
            rows.append([a, m, v[i], _rate(v[i], v[i - 1] if i else None), _rate(v[i], v[i - 12] if i >= 12 else None),
                         ev[i], ov[i], "完整" if m < cur else "进行中"])
    _table(ws, 4, ["区域", "月份", "参赛人次", "环比", "同比", "比赛场次", "超过64人比赛", "状态"], rows,
           {3: INT, 4: PCT, 5: PCT, 6: INT, 7: INT})

    wb.properties.creator = f"KaiqiuBoard {__version__}"
    wb.properties.created = datetime.now()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(snap: dict) -> str:
    return f"开球网数据_{snap['fetched_at'][:10]}.xlsx"
