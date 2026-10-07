"""python -m unittest discover -s tests -v"""
import io
import json
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from kaiqiu import excel, source, store  # noqa: E402
from kaiqiu.paths import data_dir, resource_dir  # noqa: E402
from kaiqiu.server import serve  # noqa: E402

SEED = json.loads((resource_dir() / "seed.json").read_text(encoding="utf-8"))
NOW = datetime(2026, 10, 7, 9, 0, tzinfo=store.SHANGHAI)


def monthly_csv(area, snap=SEED, bump=None):
    """按开球网的 CSV 格式还原（倒序、带 BOM、超64人可为空）。"""
    d = snap["data"][area]
    lines = ["\ufeff月份,本月比赛场次,超过64人以上的比赛场次,本月参赛人次"]
    for i in reversed(range(len(snap["months"]))):
        y, m = snap["months"][i].split("-")
        p = d["participants"][i] + (bump(area, i) if bump else 0)
        o = "" if d["over64"][i] is None else d["over64"][i]
        lines.append(f"{y}年{int(m)}月,{d['events'][i]},{o},{p}")
    return "\n".join(lines)


def period_csv(label, snap=SEED):
    rows = ["序号,省份,比赛场次,参赛人次,比赛盘数"]
    for i, r in enumerate(snap["periods"][label], 1):
        rows.append(f"{i},{r['region']},{r['events']},{r['participants']},{r['games']}")
    return "\n".join(rows)


def fake_fetch(bump=None, fail=None):
    from urllib.parse import parse_qs, urlparse

    def fetch(url):
        q = parse_qs(urlparse(url).query)
        if url.startswith(source.MONTHLY_URL):
            area = q.get("province", [source.NATIONAL])[0]
            if fail and area in fail:
                raise source.SourceError(f"连接失败 {area}")
            return monthly_csv(area, bump=bump)
        days = q.get("days", [None])[0]
        label = {None: "近7天", "30": "近30天", "90": "近90天", "365": "近365天"}[days]
        return period_csv(label)
    return fetch


def fetcher(**kw):
    return lambda progress=None: source.fetch_all(progress, fetch=fake_fetch(**kw))


class ParseTests(unittest.TestCase):
    def test_monthly_sorted_and_over64_none(self):
        m = source.parse_monthly("全国", monthly_csv("全国"))
        self.assertEqual(m.months, SEED["months"])
        self.assertEqual(m.participants, SEED["data"]["全国"]["participants"])
        self.assertIsNone(m.over64[0])

    def test_monthly_bad_header(self):
        with self.assertRaises(source.SourceError):
            source.parse_monthly("全国", "a,b,c,d\n2020年1月,1,2,3")

    def test_monthly_duplicate_month(self):
        text = "月份,本月比赛场次,超过64人以上的比赛场次,本月参赛人次\n2020年1月,1,,3\n2020年1月,1,,3"
        with self.assertRaises(source.SourceError):
            source.parse_monthly("全国", text)

    def test_period(self):
        rows = source.parse_period("近7天", period_csv("近7天"))
        self.assertEqual(len(rows), 35)
        self.assertEqual(rows[0], SEED["periods"]["近7天"][0])

    def test_urls(self):
        self.assertTrue(source.monthly_url("江苏").endswith("?province=%E6%B1%9F%E8%8B%8F"))
        self.assertEqual(source.period_url(None), source.PERIOD_URL)
        self.assertTrue(source.period_url(30).endswith("?days=30"))

    def test_mirror_fallback(self):
        calls = []
        orig = source._fetch_one
        def fake(url, retries, timeout):
            calls.append(url)
            if url.startswith(source.BASE):
                raise source.SourceError("down")
            return "ok"
        source._fetch_one = fake
        try:
            self.assertEqual(source.fetch_text(source.monthly_url("江苏")), "ok")
        finally:
            source._fetch_one = orig
        self.assertTrue(calls[1].startswith("https://kaiqiu.cc/home/cityEventsMonthly.php?province="))

    def test_fetch_all_reports_failures(self):
        with self.assertRaises(source.SourceError) as cm:
            source.fetch_all(fetch=fake_fetch(fail={"江苏", "广东"}))
        self.assertIn("江苏", str(cm.exception))


class SnapshotTests(unittest.TestCase):
    def test_build_matches_seed(self):
        monthly, periods = source.fetch_all(fetch=fake_fetch())
        snap = store.build_snapshot(monthly, periods, NOW)
        self.assertEqual(snap["months"], SEED["months"])
        self.assertEqual(snap["data"], SEED["data"])
        self.assertEqual(snap["current_month"], "2026-10")

    def test_future_month_rejected(self):
        monthly, periods = source.fetch_all(fetch=fake_fetch())
        with self.assertRaises(source.SourceError):
            store.build_snapshot(monthly, periods, datetime(2026, 9, 1, tzinfo=store.SHANGHAI))

    def test_diff_ignores_open_month_growth(self):
        old = json.loads(json.dumps(SEED)); old["seed"] = False
        new = json.loads(json.dumps(SEED))
        new["data"]["江苏"]["participants"][-1] += 100   # 进行中月份：正常增长
        new["data"]["江苏"]["participants"][100] += 7    # 历史月份：网站修订
        d = store.diff(old, new)
        self.assertEqual(d["revised_count"], 1)
        self.assertEqual(d["revised"][0]["month"], SEED["months"][100])

    def test_sanity_rejects_collapse(self):
        old = json.loads(json.dumps(SEED)); old.pop("seed")
        new = json.loads(json.dumps(SEED))
        new["data"]["全国"]["participants"] = [v // 2 for v in new["data"]["全国"]["participants"]]
        with self.assertRaises(source.SourceError):
            store.sanity_check(old, new)
        store.sanity_check(SEED, new)  # 附带数据不作为比较基准


class StoreTests(unittest.TestCase):
    def wait(self, st):
        for _ in range(200):
            if not st.progress["running"]:
                return
            time.sleep(0.02)
        self.fail("刷新没有结束")

    def test_refresh_success_then_failure_keeps_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            st = store.Store(Path(tmp), fetcher=fetcher())
            self.assertTrue(st.snapshot["seed"])
            self.assertTrue(st.should_refresh_on_open())
            self.assertTrue(st.start_refresh()); self.wait(st)
            self.assertIsNone(st.progress["error"], st.progress)
            rev = st.snapshot["revision"]
            self.assertFalse(st.should_refresh_on_open())  # 刚更新过，不重复抓
            st.fetcher = fetcher(fail={"海外"})
            st.start_refresh(); self.wait(st)
            self.assertIn("海外", st.progress["error"])
            self.assertEqual(st.snapshot["revision"], rev)
            # 重启后恢复同一版本
            st2 = store.Store(Path(tmp))
            self.assertEqual(st2.snapshot["revision"], rev)

    def test_revision_detected_and_pruned(self):
        with tempfile.TemporaryDirectory() as tmp:
            st = store.Store(Path(tmp), fetcher=fetcher())
            st.start_refresh(); self.wait(st)
            st.fetcher = fetcher(bump=lambda a, i: 5 if (a == "北京" and i == 150) else 0)
            st.start_refresh(); self.wait(st)
            self.assertEqual(st.snapshot["changes"]["revised_count"], 1)
            for _ in range(store.KEEP_SNAPSHOTS + 2):
                st.start_refresh(); self.wait(st)
            self.assertLessEqual(len(list(st.snap_dir.glob("*.json"))), store.KEEP_SNAPSHOTS)

    def test_settings_validated_and_persisted(self):
        with tempfile.TemporaryDirectory() as tmp:
            st = store.Store(Path(tmp))
            st.set_settings({"periodic_hours": 6, "auto_refresh_on_open": False, "min_refresh_minutes": "x"})
            self.assertEqual(store.Store(Path(tmp)).settings,
                             {"auto_refresh_on_open": False, "min_refresh_minutes": 30, "periodic_hours": 6})
            st.set_settings({"periodic_hours": 5})
            self.assertEqual(st.settings["periodic_hours"], 6)

    def test_corrupt_pointer_falls_back_to_seed(self):
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, "current.json").write_text('{"revision":"../../etc"}', encoding="utf-8")
            self.assertTrue(store.Store(Path(tmp)).snapshot["seed"])


class ExcelTests(unittest.TestCase):
    def test_workbook(self):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(excel.build(SEED)))
        self.assertEqual(wb.sheetnames, ["说明", "全国月度", "区域排名", "参赛人次宽表", "比赛场次宽表", "年度汇总", "时段统计", "区域月度明细"])
        ws = wb["全国月度"]
        self.assertEqual(ws.cell(5, 1).value, "2010-01")
        self.assertEqual(ws.cell(4 + len(SEED["months"]), 2).value, SEED["data"]["全国"]["participants"][-1])
        rank = wb["区域排名"]
        self.assertEqual(rank.cell(5, 2).value, "广东")
        self.assertEqual(wb["区域月度明细"].max_row, 4 + 36 * len(SEED["months"]))
        # 不写公式，全部是数值
        self.assertFalse(any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row))
        self.assertEqual(len(wb["全国月度"]._charts) + len(rank._charts) + len(wb["年度汇总"]._charts), 3)


class PathTests(unittest.TestCase):
    def test_data_dir(self):
        self.assertEqual(data_dir("win32", {"LOCALAPPDATA": r"C:\Users\a\AppData\Local"}, Path("/h")),
                         Path(r"C:\Users\a\AppData\Local") / "KaiqiuBoard")
        self.assertEqual(data_dir("darwin", {}, Path("/Users/a")), Path("/Users/a/Library/Application Support/KaiqiuBoard"))
        self.assertEqual(data_dir("linux", {"KAIQIU_BOARD_DATA": "/x"}, Path("/h")), Path("/x"))


class ServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.store = store.Store(Path(cls.tmp.name), fetcher=fetcher())
        cls.app, cls.httpd = serve(cls.store)
        cls.port = cls.httpd.server_address[1]
        threading.Thread(target=cls.httpd.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown(); cls.httpd.server_close(); cls.tmp.cleanup()

    def req(self, path, body=None, headers=None):
        h = {"Content-Type": "application/json", "X-Kaiqiu-Token": self.app.token} if body is not None else {}
        h.update(headers or {})
        r = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", data=None if body is None else json.dumps(body).encode(), headers=h)
        try:
            with urllib.request.urlopen(r) as resp:
                return resp.status, resp.headers, resp.read()
        except urllib.error.HTTPError as e:
            return e.code, e.headers, e.read()

    def test_static_and_data(self):
        for p in ("/", "/app.js", "/model.js", "/style.css", "/icon.svg", "/vendor/echarts.min.js"):
            self.assertEqual(self.req(p)[0], 200, p)
        code, headers, body = self.req("/api/data")
        self.assertEqual(code, 200)
        self.assertEqual(json.loads(body)["months"][0], "2010-01")
        self.assertEqual(self.req("/api/data", headers={"If-None-Match": headers["ETag"]})[0], 304)
        self.assertIn("frame-ancestors 'none'", headers["Content-Security-Policy"])

    def test_rejects_foreign_host_origin_and_token(self):
        self.assertEqual(self.req("/api/status", headers={"Host": "evil.example"})[0], 403)
        self.assertEqual(self.req("/api/refresh", {}, {"X-Kaiqiu-Token": "wrong"})[0], 403)
        self.assertEqual(self.req("/api/refresh", {}, {"Origin": "http://evil.example"})[0], 403)
        self.assertEqual(self.req("/api/reveal", {"path": "/etc/passwd"})[0], 403)

    def test_export_and_save_browser_mode(self):
        code, headers, body = self.req("/api/export.xlsx")
        self.assertEqual(code, 200)
        self.assertTrue(body.startswith(b"PK"))
        self.assertIn("filename*=UTF-8''", headers["Content-Disposition"])
        code, _, body = self.req("/api/save", {"kind": "xlsx"})
        self.assertEqual(json.loads(body), {"mode": "browser"})

    def test_save_desktop_mode(self):
        class FakeDesktop:
            def __init__(self, target): self.target = target
            def save_dialog(self, name, types): return self.target
        with tempfile.TemporaryDirectory() as tmp:
            target = str(Path(tmp, "out.csv"))
            self.app.desktop = FakeDesktop(target)
            try:
                code, _, body = self.req("/api/save", {"kind": "csv", "name": "x.csv", "content": "a,b\n1,2"})
                self.assertEqual(json.loads(body), {"saved": target})
                self.assertEqual(Path(target).read_bytes(), "\ufeffa,b\n1,2".encode())
            finally:
                self.app.desktop = None

    def test_refresh_endpoint(self):
        code, _, body = self.req("/api/refresh", {})
        self.assertEqual(code, 202)
        for _ in range(200):
            if not self.store.progress["running"]:
                break
            time.sleep(0.02)
        self.assertIsNone(self.store.progress["error"])
        self.assertFalse(json.loads(self.req("/api/status")[2])["seed"])


if __name__ == "__main__":
    unittest.main()
